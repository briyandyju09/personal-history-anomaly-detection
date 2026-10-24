# -*- coding: utf-8 -*-
"""Build presentation.pptx from slides_spec.json (same content as main.tex).

    pip install python-pptx
    python generate_pptx.py
"""
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
spec = json.loads((HERE / "slides_spec.json").read_text(encoding="utf-8"))
FIG = HERE / "figures"

PRIMARY = RGBColor(0x2D, 0x5F, 0x8B)
DEEP = RGBColor(0x1F, 0x45, 0x66)
GOLD = RGBColor(0xC9, 0xA2, 0x4A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x88, 0x88, 0x88)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = len(spec["slides"])


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def add_runs(par, text, size, color=INK, bold=False):
    """Write text into a paragraph, honouring **bold** markers."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run()
        r.text = chunk
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.bold = bold or (i % 2 == 1)
        r.font.color.rgb = color


def title_bar(slide, text):
    tf = textbox(slide, 0.6, 0.35, 12.1, 1.0, MSO_ANCHOR.MIDDLE)
    add_runs(tf.paragraphs[0], text, 38, PRIMARY, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.38), Inches(12.1), Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = GOLD
    rule.line.fill.background()


def footer(slide, n):
    tf = textbox(slide, 0.6, 6.95, 9, 0.4)
    add_runs(tf.paragraphs[0], "Dyju · How Much Personal History Is Enough?", 12, MUTED)
    tf = textbox(slide, 11.2, 6.95, 1.5, 0.4)
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_runs(tf.paragraphs[0], f"{n}/{TOTAL}", 12, MUTED)


def bullets(slide, items, x, y, w, h, size):
    tf = textbox(slide, x, y, w, h)
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.55)
        add_runs(p, "■  ", size * 0.7, PRIMARY)
        add_runs(p, b, size)


def notes(slide, s):
    t = f"{s['say']} [{s['time'] // 60}:{s['time'] % 60:02d}]"
    if s.get("transition"):
        t += f" Transition: {s['transition']}"
    slide.notes_slide.notes_text_frame.text = t


def picture(slide, name, x, y, max_w, max_h):
    from PIL import Image
    path = FIG / f"{name}.png"
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def table(slide, tb, x, y, w, size):
    rows, cols = len(tb["rows"]) + 1, len(tb["header"])
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(0.55 * rows))
    t = shape.table
    for c in range(cols):
        t.columns[c].width = Inches(w / cols)
    for r in range(rows):
        for c in range(cols):
            cell = t.cell(r, c)
            txt = tb["header"][c] if r == 0 else tb["rows"][r - 1][c]
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            add_runs(p, txt, size, RGBColor(255, 255, 255) if r == 0 else INK, bold=(r == 0))
            cell.fill.solid()
            cell.fill.fore_color.rgb = PRIMARY if r == 0 else (RGBColor(0xE8, 0xF1, 0xF8) if r % 2 else RGBColor(255, 255, 255))


for s in spec["slides"]:
    slide = prs.slides.add_slide(BLANK)
    t = s["type"]
    if t == "title":
        tf = textbox(slide, 0.9, 1.5, 11.5, 2.0, MSO_ANCHOR.MIDDLE)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], spec["title"], 50, DEEP, bold=True)
        tf = textbox(slide, 1.2, 3.6, 10.9, 1.0)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], "Measuring the data requirement for individualised financial behavioural anomaly detection", 24, PRIMARY)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.4), Inches(4.75), Inches(2.5), Inches(0.05))
        rule.fill.solid()
        rule.fill.fore_color.rgb = GOLD
        rule.line.fill.background()
        tf = textbox(slide, 1.2, 5.0, 10.9, 1.6)
        for i, (txt, sz, col) in enumerate([(spec["author"], 28, INK), (spec["institute"], 20, MUTED),
                                            ("Independent research, 2026", 18, MUTED)]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            add_runs(p, txt, sz, col)
        notes(slide, s)
        continue

    title_bar(slide, s["title"])
    if t == "bullets":
        y = 1.8
        bullets(slide, s["bullets"], 0.8, y, 11.7, 3.8, 30)
        if s.get("highlight"):
            tf = textbox(slide, 0.8, 5.3, 11.7, 1.0, MSO_ANCHOR.MIDDLE)
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            add_runs(tf.paragraphs[0], s["highlight"], 36, PRIMARY, bold=True)
    elif t == "fig":
        picture(slide, s["fig"], 0.6, 1.75, 8.3, 4.9)
        bullets(slide, s["bullets"], 9.1, 2.2, 3.8, 4.2, 24)
    elif t == "table":
        tb = s["table"]
        if len(tb["header"]) > 3:
            table(slide, tb, 1.3, 1.8, 10.7, 24)
            bullets(slide, s["bullets"], 0.8, 4.9, 11.7, 1.9, 26)
        else:
            table(slide, tb, 0.8, 2.0, 6.2, 24)
            bullets(slide, s["bullets"], 7.4, 2.0, 5.4, 3.5, 26)
    elif t == "eq":
        tf = textbox(slide, 0.8, 1.7, 11.7, 1.3, MSO_ANCHOR.MIDDLE)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], s["eq_txt"], 40, DEEP, bold=True)
        bullets(slide, s["bullets"], 0.8, 3.4, 11.7, 3.2, 28)
    elif t == "qr":
        slide.shapes.add_picture(str(FIG / "qr_paper.png"), Inches(1.4), Inches(2.0), Inches(4.0), Inches(4.0))
        tf = textbox(slide, 6.3, 2.5, 6.4, 3.2)
        add_runs(tf.paragraphs[0], "Paper and code", 34, PRIMARY, bold=True)
        for b in s["bullets"]:
            p = tf.add_paragraph()
            p.space_before = Pt(14)
            add_runs(p, b, 28)
    footer(slide, s["id"])
    notes(slide, s)

out = HERE / "presentation.pptx"
prs.save(out)
print("wrote", out, "-", len(prs.slides), "slides")
